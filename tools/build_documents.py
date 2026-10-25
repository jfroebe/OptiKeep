"""Build versioned engineering DOCX files from their Markdown sources.

Copyright (C) 2026 Jason Froebe. GPL-3.0-only.
Requires python-docx; PDF conversion is performed by the document renderer.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build(source: Path):
    """Render prose and the compact requirements matrix with stable page breaks."""
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(22)
    document.styles["Heading 1"].font.size = Pt(18)
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.text = "OPTIKEEP   /   ENGINEERING SPECIFICATION   /   1.1"
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.text = "OptiKeep   |   25 October 2026   |   "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    first = True
    table = None
    for line in source.read_text().splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            if not first:
                document.add_page_break()
            if first:
                document.add_paragraph(source.stem.replace("_", " "), "Title")
            document.add_heading(line[2:], 1)
            first = False
            table = None
        elif line == "Requirements implementation matrix":
            document.add_page_break()
            document.add_heading(line, 1)
        elif line.startswith("|"):
            values = [cell.strip() for cell in line.strip("|").split("|")]
            if all(set(value) <= {"-", " "} for value in values):
                continue
            if table is None:
                table = document.add_table(rows=0, cols=len(values))
                table.style = "Table Grid"
            cells = table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.space_before = Pt(5)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                properties.append(borders)
                if len(table.rows) == 1:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7EEF2")
                    properties.append(shade)
        else:
            paragraph = document.add_paragraph(line.replace(" -> ", " then "))
            if len(line) < 95 and not line.startswith("["):
                paragraph.runs[0].bold = True
    document.core_properties.author = "Jason Froebe"
    document.core_properties.title = source.stem.replace("_", " ")
    document.save(source.with_suffix(".docx"))


if __name__ == "__main__":
    for path in (Path(__file__).resolve().parents[1] / "docs").glob("*/*.md"):
        build(path)
